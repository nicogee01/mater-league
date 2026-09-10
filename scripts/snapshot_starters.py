"""
Snapshot Weekly Starters
------------------------
Runs on a schedule via GitHub Actions. Each run:

1. Pulls current rosters from Sleeper.
2. Compares each team's (wins + losses + ties) to the total saved from the
   previous run (stored in data/last_totals.json, committed to the repo).
3. If a team's total went UP since last run, that team just had a result
   locked in — so we grab their CURRENT full roster (starters AND bench,
   before anyone can change it for the next matchweek) and record it as
   that week's roster snapshot, with an explicit Starter flag per player.
4. Everything gets appended to starters.xlsx as (Week, Manager, Player,
   Starter), matching the manually-entered weeks 1-3.

Recording the full roster (not just starters) matters because a player who
gets dropped later in the season would otherwise disappear from the site's
current-roster view entirely — this way their bench appearance from a past
week stays on record even after they're no longer on anyone's roster.

Safe to re-run: it never overwrites a (week, manager) pair that's already
recorded, and the very first run just establishes a baseline without
recording anything (since there's nothing to compare against yet).
"""

import json
import os
import requests
from openpyxl import Workbook, load_workbook

LEAGUE_ID = "1390750210698780672"
SPORT = "clubsoccer:epl"
API = "https://api.sleeper.app/v1"
STARTERS_FILE = "starters.xlsx"
STATE_FILE = "data/last_totals.json"


def get_json(url):
    r = requests.get(url, timeout=30)
    r.raise_for_status()
    return r.json()


def manager_name(roster, users_by_id):
    u = users_by_id.get(roster.get("owner_id"))
    return u["display_name"] if u else "Unknown"


def player_name(pid, players):
    p = players.get(pid)
    if not p:
        return f"Player {pid}"
    name = p.get("full_name")
    if not name:
        name = f"{p.get('first_name', '')} {p.get('last_name', '')}".strip()
    return name or f"Player {pid}"


def main():
    rosters = get_json(f"{API}/league/{LEAGUE_ID}/rosters")
    users = get_json(f"{API}/league/{LEAGUE_ID}/users")
    state = get_json(f"{API}/state/{SPORT}")

    # The week whose results just finished (not the upcoming/in-progress week)
    week_to_record = state.get("last_scored_leg") or (state.get("week", 1) - 1)

    os.makedirs("data", exist_ok=True)
    last_totals = {}
    if os.path.exists(STATE_FILE):
        with open(STATE_FILE) as f:
            last_totals = json.load(f)

    users_by_id = {u["user_id"]: u for u in users}

    new_totals = {}
    changed_rosters = []
    for r in rosters:
        rid = str(r["roster_id"])
        settings = r.get("settings", {})
        total = (settings.get("wins", 0) or 0) + (settings.get("losses", 0) or 0) + (settings.get("ties", 0) or 0)
        new_totals[rid] = total
        prev = last_totals.get(rid)
        if prev is not None and total > prev:
            changed_rosters.append(r)

    if week_to_record < 1 or not changed_rosters:
        print(f"Nothing to record this run (week_to_record={week_to_record}, changed={len(changed_rosters)}).")
        with open(STATE_FILE, "w") as f:
            json.dump(new_totals, f)
        return

    players = get_json(f"{API}/players/{SPORT}")

    if os.path.exists(STARTERS_FILE):
        wb = load_workbook(STARTERS_FILE)
        ws = wb["Starters"] if "Starters" in wb.sheetnames else wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Starters"
        ws.append(["Week", "Manager", "Player", "Starter"])

    existing = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        existing.add((int(row[0]), str(row[1]).strip().lower()))

    recorded_any = False
    for r in changed_rosters:
        mgr = manager_name(r, users_by_id)
        key = (week_to_record, mgr.strip().lower())
        if key in existing:
            continue
        starter_set = set(r.get("starters", []))
        for pid in r.get("players", []):
            is_starter = 1 if pid in starter_set else 0
            ws.append([week_to_record, mgr, player_name(pid, players), is_starter])
        print(f"Recorded full roster for {mgr}, week {week_to_record}")
        recorded_any = True

    if recorded_any:
        wb.save(STARTERS_FILE)

    with open(STATE_FILE, "w") as f:
        json.dump(new_totals, f)


if __name__ == "__main__":
    main()
